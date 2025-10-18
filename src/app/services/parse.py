"""Document parsing service for various file formats"""
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class DocumentParser:
    """Parse documents of various formats into text"""

    @staticmethod
    def parse(file_bytes: bytes, mime_type: str, filename: str) -> str:
        """
        Parse document bytes into text based on MIME type

        Args:
            file_bytes: Raw file bytes
            mime_type: MIME type of the file
            filename: Original filename

        Returns:
            Extracted text content

        Raises:
            ValueError: If file format is not supported
        """
        if mime_type == "application/pdf":
            return DocumentParser._parse_pdf(file_bytes)
        elif mime_type in [
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        ]:
            return DocumentParser._parse_docx(file_bytes)
        elif mime_type == "text/plain":
            return DocumentParser._parse_text(file_bytes)
        elif mime_type == "text/markdown":
            return DocumentParser._parse_text(file_bytes)
        elif mime_type == "text/csv":
            return DocumentParser._parse_text(file_bytes)
        else:
            raise ValueError(f"Unsupported file format: {mime_type}")

    @staticmethod
    def _parse_pdf(file_bytes: bytes) -> str:
        """Parse PDF file"""
        try:
            from pypdf import PdfReader
            from io import BytesIO

            pdf = PdfReader(BytesIO(file_bytes))
            text_parts = []

            for page_num, page in enumerate(pdf.pages):
                text = page.extract_text()
                if text.strip():
                    text_parts.append(text)

            full_text = "\n\n".join(text_parts)
            logger.info(f"Extracted {len(full_text)} chars from PDF with {len(pdf.pages)} pages")
            return full_text

        except Exception as e:
            logger.error(f"Failed to parse PDF: {e}")
            raise ValueError(f"Failed to parse PDF: {str(e)}")

    @staticmethod
    def _parse_docx(file_bytes: bytes) -> str:
        """Parse DOCX file"""
        try:
            from docx import Document
            from io import BytesIO

            doc = Document(BytesIO(file_bytes))
            text_parts = []

            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)

            full_text = "\n\n".join(text_parts)
            logger.info(f"Extracted {len(full_text)} chars from DOCX")
            return full_text

        except Exception as e:
            logger.error(f"Failed to parse DOCX: {e}")
            raise ValueError(f"Failed to parse DOCX: {str(e)}")

    @staticmethod
    def _parse_text(file_bytes: bytes) -> str:
        """Parse plain text file"""
        try:
            # Try UTF-8 first
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            # Fall back to latin-1
            try:
                text = file_bytes.decode("latin-1")
            except Exception as e:
                logger.error(f"Failed to decode text file: {e}")
                raise ValueError(f"Failed to decode text file: {str(e)}")

        logger.info(f"Extracted {len(text)} chars from text file")
        return text


# Singleton instance
document_parser = DocumentParser()
